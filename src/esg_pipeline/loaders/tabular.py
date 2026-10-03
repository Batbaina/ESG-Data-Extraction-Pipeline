from pathlib import Path
from openpyxl import load_workbook
from .text import make_records

def load_tabular(path: Path, text_column: str = "text", limit: int | None = None):
    suffix = path.suffix.lower()
    rows = []
    if suffix == ".xlsx":
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        headers = [str(v).strip() if v is not None else "" for v in next(ws.iter_rows(values_only=True))]
        if text_column not in headers:
            raise ValueError(f"Column '{text_column}' not found. Available: {headers}")
        idx = headers.index(text_column)
        for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            val = row[idx] if idx < len(row) else None
            if val and str(val).strip():
                meta = {headers[j]: row[j] for j in range(min(len(headers), len(row))) if headers[j] and j != idx}
                rows.append({"text": str(val).strip(), "source_locator": f"row:{i}", "metadata": meta})
                if limit and len(rows) >= limit:
                    break
        return rows
    if suffix == ".csv":
        import csv
        with path.open("r", encoding="utf-8-sig", newline="") as f:
            reader = csv.DictReader(f)
            if text_column not in (reader.fieldnames or []):
                raise ValueError(f"Column '{text_column}' not found. Available: {reader.fieldnames}")
            for i, row in enumerate(reader, start=2):
                val = row.get(text_column)
                if val and val.strip():
                    meta = {k:v for k,v in row.items() if k != text_column}
                    rows.append({"text": val.strip(), "source_locator": f"row:{i}", "metadata": meta})
                    if limit and len(rows) >= limit:
                        break
        return rows
    raise ValueError("Supported tabular formats: .xlsx, .csv")
