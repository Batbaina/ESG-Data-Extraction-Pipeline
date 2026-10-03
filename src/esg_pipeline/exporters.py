import json
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HEADERS = [
    "source_id","source_file","source_locator","source_text","status","usable","rejection_reason",
    "pillar","topic","knowledge_type","source_facts","extracted_knowledge","confidence",
    "qc_pass","qc_lexical_grounding","qc_reasons","llm_retry_count","model","prompt_version",
    "temperature","top_p","top_k","seed","num_ctx","num_predict",
    "human_esg_relevance","human_faithfulness","human_completeness",
    "human_self_contained","human_training_usefulness","human_total","human_comment"
]


def write_jsonl(path: Path, rows):
    with path.open("w", encoding="utf-8") as f:
        for row in rows:
            f.write(json.dumps(row, ensure_ascii=False) + "\n")


def write_json(path: Path, obj):
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding="utf-8")


def write_review_xlsx(path: Path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "human_review"
    ws.append(HEADERS)
    for row in rows:
        ws.append([row.get(h, "") for h in HEADERS])

    header_fill = PatternFill("solid", fgColor="173F5F")
    review_fill = PatternFill("solid", fgColor="FCE8B2")
    imported_font = "008000"
    review_start = HEADERS.index("human_esg_relevance") + 1
    review_end = HEADERS.index("human_comment") + 1
    total_col = HEADERS.index("human_total") + 1

    for c in ws[1]:
        c.fill = header_fill
        c.font = Font(color="FFFFFF", bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="center")

    for r in range(2, ws.max_row + 1):
        for col in range(1, review_start):
            ws.cell(r, col).font = Font(color=imported_font)
        for col in range(review_start, review_end + 1):
            ws.cell(r, col).fill = review_fill
        # five human dimensions are immediately before total/comment
        first_score = get_column_letter(review_start)
        last_score = get_column_letter(review_start + 4)
        ws.cell(r, total_col, f'=IF(COUNT({first_score}{r}:{last_score}{r})=5,SUM({first_score}{r}:{last_score}{r}),"")')

    widths = {
        "source_id":18,"source_file":24,"source_locator":20,"source_text":70,"status":14,"usable":10,
        "rejection_reason":28,"pillar":16,"topic":24,"knowledge_type":18,"source_facts":55,
        "extracted_knowledge":75,"confidence":12,"qc_pass":12,"qc_lexical_grounding":18,"qc_reasons":28,
        "llm_retry_count":14,"model":18,"prompt_version":14,"temperature":12,"top_p":10,"top_k":10,
        "seed":10,"num_ctx":12,"num_predict":12,"human_esg_relevance":18,"human_faithfulness":18,
        "human_completeness":18,"human_self_contained":20,"human_training_usefulness":22,"human_total":12,
        "human_comment":40,
    }
    for i, h in enumerate(HEADERS, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(h, 16)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)


def finalize_review(review_path: Path, output_jsonl: Path, min_human_score: int = 8):
    wb = load_workbook(review_path, data_only=False)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    score_cols = [
        "human_esg_relevance","human_faithfulness","human_completeness",
        "human_self_contained","human_training_usefulness"
    ]
    final = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        row = {headers[i]: values[i] for i in range(len(headers))}
        scores = [row.get(c) for c in score_cols]
        if not all(isinstance(v, (int, float)) and 0 <= v <= 2 for v in scores):
            continue
        total = sum(scores)
        if row.get("qc_pass") is True and total >= min_human_score and row.get("extracted_knowledge"):
            final.append({
                "text": row["extracted_knowledge"],
                "pillar": row["pillar"],
                "topic": row["topic"],
                "knowledge_type": row["knowledge_type"],
                "source_id": row["source_id"],
                "source_file": row["source_file"],
                "source_locator": row["source_locator"],
                "human_score": total,
                "generation_metadata": {
                    "model": row.get("model"), "prompt_version": row.get("prompt_version"),
                    "temperature": row.get("temperature"), "top_p": row.get("top_p"),
                    "top_k": row.get("top_k"), "seed": row.get("seed"),
                    "num_ctx": row.get("num_ctx"), "num_predict": row.get("num_predict"),
                },
            })
    write_jsonl(output_jsonl, final)
    return len(final)
